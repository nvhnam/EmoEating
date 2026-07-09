"""
EmoEating — Plain User Flow Diagram (PPTX for Canva)
Output: paper/figures/emoeating_userflow.pptx

One tall portrait slide, black-and-white only.
Boxes with plain text + arrows between steps.
No colors, no emojis, no title, no header, no footer.
Upload to Canva → Import → ungroup to edit individual shapes.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
FONT  = "Arial"

# ── Slide dimensions: 8 × 22 inch tall portrait ────────────────────────────
W = Inches(8)
H = Inches(22)

# Layout
MARGIN  = 0.3     # inches, left/right margin
BOX_W   = 8 - MARGIN * 2    # 7.4" main box
FORK_W  = 3.3                # width of each side-by-side fork box
FORK_GAP = 8 - MARGIN * 2 - FORK_W * 2   # gap between fork boxes = 0.8"
FORK_LX = MARGIN             # left fork box x
FORK_RX = MARGIN + FORK_W + FORK_GAP    # right fork box x
CX      = 8 / 2              # center x = 4.0"
FLC     = FORK_LX + FORK_W / 2   # left fork center  = 1.95"
FRC     = FORK_RX + FORK_W / 2   # right fork center = 6.05"


def make_slide(prs):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    return s


def add_box(slide, x, y, w, h, lines, dashed=False, font_size=11):
    """Rounded rectangle with centered multi-line text, black border, white fill."""
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    # Rounded corners
    sp = shape.element
    pg = sp.find('.//' + qn('a:prstGeom'))
    if pg is not None:
        pg.set('prst', 'roundRect')
        avLst = pg.find(qn('a:avLst'))
        if avLst is None:
            avLst = etree.SubElement(pg, qn('a:avLst'))
        else:
            for child in list(avLst):
                avLst.remove(child)
        gd = etree.SubElement(avLst, qn('a:gd'))
        gd.set('name', 'adj')
        gd.set('fmla', 'val 16667')

    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE

    shape.line.color.rgb = BLACK
    shape.line.width = Pt(1.2)

    if dashed:
        ln_el = sp.find('.//' + qn('a:ln'))
        if ln_el is not None:
            pd = etree.SubElement(ln_el, qn('a:prstDash'))
            pd.set('val', 'dash')

    tf = shape.text_frame
    tf.word_wrap = True
    tf._txBody.set('anchor', 'ctr')
    tf.margin_top    = Pt(6)
    tf.margin_bottom = Pt(6)
    tf.margin_left   = Pt(8)
    tf.margin_right  = Pt(8)

    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(font_size)
        run.font.color.rgb = BLACK
        run.font.bold = False
    return shape


def add_vline(slide, x, y1, y2):
    """Plain vertical line, no arrowhead."""
    c = slide.shapes.add_connector(1, Inches(x), Inches(y1), Inches(x), Inches(y2))
    c.line.color.rgb = BLACK
    c.line.width = Pt(1.2)
    return c


def add_hline(slide, x1, x2, y):
    """Plain horizontal line."""
    c = slide.shapes.add_connector(1, Inches(x1), Inches(y), Inches(x2), Inches(y))
    c.line.color.rgb = BLACK
    c.line.width = Pt(1.2)
    return c


def add_arrow_down(slide, x, y1, y2):
    """Vertical arrow pointing downward (arrowhead at bottom)."""
    c = slide.shapes.add_connector(1, Inches(x), Inches(y1), Inches(x), Inches(y2))
    c.line.color.rgb = BLACK
    c.line.width = Pt(1.2)
    # Arrowhead at tail end (the lower point)
    sp = c.element
    ln_el = sp.find('.//' + qn('a:ln'))
    if ln_el is not None:
        # Remove any existing end markers
        for tag in (qn('a:headEnd'), qn('a:tailEnd')):
            el = ln_el.find(tag)
            if el is not None:
                ln_el.remove(el)
        tail = etree.SubElement(ln_el, qn('a:tailEnd'))
        tail.set('type', 'arrow')
        tail.set('w', 'med')
        tail.set('len', 'med')
    return c


def box_h(n_lines, font_size=11):
    line_h = font_size * 1.4 / 72   # approx inches per line
    pad = 0.18
    return n_lines * line_h + pad


def build():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    slide = make_slide(prs)
    y = 0.30   # cursor: current Y position

    ARROW = 0.26   # height of each arrow gap

    # ─── BOX 1: Consent ────────────────────────────────────────────────────
    b1 = [
        "Read research ethics notice",
        "Tick consent checkbox",
        "Click Start",
    ]
    h1 = box_h(len(b1))
    add_box(slide, MARGIN, y, BOX_W, h1, b1)
    y += h1

    add_arrow_down(slide, CX, y, y + ARROW)
    y += ARROW

    # ─── BOX 2: Profile (optional, dashed) ────────────────────────────────
    b2 = [
        "Enter: age, sex, height, weight, activity level, meal type, dietary restrictions",
        "System computes: BMI  |  BMR  |  TDEE  |  meal calorie target per meal type",
        "(Step is optional — user may skip to next screen)",
    ]
    h2 = box_h(len(b2))
    add_box(slide, MARGIN, y, BOX_W, h2, b2, dashed=True)
    y += h2

    # ─── FORK: stem → horizontal split ────────────────────────────────────
    STEM = 0.20
    add_vline(slide, CX, y, y + STEM)
    y += STEM
    add_hline(slide, FLC, FRC, y)

    FORK_DROP = 0.18
    add_arrow_down(slide, FLC, y, y + FORK_DROP)
    add_arrow_down(slide, FRC, y, y + FORK_DROP)
    y += FORK_DROP

    # ─── BOX 3L + 3R: Voice / Manual ──────────────────────────────────────
    b3l = [
        "Voice Input  (Primary Path)",
        "",
        "Record 3-5 seconds of natural speech",
        'Prompt: "Kids are talking by the door."',
        "SER model: emotion2vec_plus_large",
        "Outputs 9-class emotion probabilities",
        "Highest confidence class -> emotional zone",
    ]
    b3r = [
        "Manual Selection  (Alternative Path)",
        "",
        "Choose from 11-emotion affect grid:",
        "Happy, Excited, Calm, Peaceful, Neutral,",
        "Anxious, Stressed, Angry, Frustrated, Sad, Tired",
        "",
        "Selection maps to valence/arousal coordinates",
    ]
    h3 = box_h(max(len(b3l), len(b3r)))
    add_box(slide, FORK_LX, y, FORK_W, h3, b3l)
    add_box(slide, FORK_RX, y, FORK_W, h3, b3r)
    y += h3

    # ─── MERGE: two lines down + horizontal + arrow down ──────────────────
    MERGE_DROP = 0.20
    add_vline(slide, FLC, y, y + MERGE_DROP)
    add_vline(slide, FRC, y, y + MERGE_DROP)
    y += MERGE_DROP
    add_hline(slide, FLC, FRC, y)
    add_arrow_down(slide, CX, y, y + FORK_DROP)
    y += FORK_DROP

    # ─── BOX 4: Zone Classification ────────────────────────────────────────
    b4 = [
        "Emotion placed on Russell Circumplex (Valence x Arousal plane)",
        "Classified into one of four affective zones:",
        "Q1 Positive-Active  |  Q2 Negative-Active",
        "Q3 Negative-Deactive  |  Q4 Neutral/Calm",
        "User rates confidence in detected emotion (scale 1-5)",
        "Click: Get Recommendations",
    ]
    h4 = box_h(len(b4))
    add_box(slide, MARGIN, y, BOX_W, h4, b4)
    y += h4

    add_arrow_down(slide, CX, y, y + ARROW)
    y += ARROW

    # ─── BOX 5: Need Vector ────────────────────────────────────────────────
    b5 = [
        "System builds Nutritional Need Vector for the detected zone:",
        "Zone-specific macro ratio targets: carbohydrate, protein, fat (percentages and grams)",
        "Zone-specific priority micronutrients: vitamins and minerals grounded in cited literature",
    ]
    h5 = box_h(len(b5))
    add_box(slide, MARGIN, y, BOX_W, h5, b5)
    y += h5

    add_arrow_down(slide, CX, y, y + ARROW)
    y += ARROW

    # ─── BOX 6: ENMS Scoring ───────────────────────────────────────────────
    b6 = [
        "Each food in database scored by Emotion-Nutrition Matching Score (ENMS):",
        "ENMS  =  0.55 x macro_score  +  0.20 x micro_score  +  0.25 x preference_score",
        "macro_score: how well carb/protein/fat match zone targets",
        "micro_score: coverage of zone priority vitamins and minerals",
        "preference_score: user history (initialized neutral at 0.5)",
    ]
    h6 = box_h(len(b6))
    add_box(slide, MARGIN, y, BOX_W, h6, b6)
    y += h6

    add_arrow_down(slide, CX, y, y + ARROW)
    y += ARROW

    # ─── BOX 7: Recommendations ────────────────────────────────────────────
    b7 = [
        "Top 3 meals displayed, ranked by ENMS score",
        "Each meal card: name, ENMS score, calorie count vs. personal target",
        "Macro fulfillment bars (carb/protein/fat)",
        "Priority micronutrient coverage, dietary flags, prep time estimate",
        "Expandable: full nutrient table, ingredient list, food photo gallery",
    ]
    h7 = box_h(len(b7))
    add_box(slide, MARGIN, y, BOX_W, h7, b7)
    y += h7

    add_arrow_down(slide, CX, y, y + ARROW)
    y += ARROW

    # ─── BOX 8: Restaurant Finder (optional, dashed) ───────────────────────
    b8 = [
        "Enter city name or allow location access (GPS)",
        "Finds restaurants serving recommended dishes within 11 km radius",
        "(Step is optional)",
    ]
    h8 = box_h(len(b8))
    add_box(slide, MARGIN, y, BOX_W, h8, b8, dashed=True)
    y += h8

    add_arrow_down(slide, CX, y, y + ARROW)
    y += ARROW

    # ─── BOX 9: User Selects Meal ──────────────────────────────────────────
    b9 = [
        'User taps "I ate it" on chosen meal',
        "Anonymous session data logged to research database:",
        "emotion, affective zone, ENMS scores, food item selected",
    ]
    h9 = box_h(len(b9))
    add_box(slide, MARGIN, y, BOX_W, h9, b9)
    y += h9

    # Resize slide height to fit content + bottom margin
    actual_h = Inches(y + 0.35)
    prs.slide_height = actual_h

    out_dir = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        'paper', 'figures'
    )
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, 'emoeating_userflow.pptx')
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == '__main__':
    build()
